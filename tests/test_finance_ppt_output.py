import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

from tools import finance_engine  # noqa: E402
from tools.export_ppt import generate_ppt  # noqa: E402


OUTPUT_DIR = ROOT / "validation_outputs" / "finance_analysis_stub"
OUTPUT_PPT = OUTPUT_DIR / "finance_analysis_stub.pptx"


def _finance_payload():
    frame = finance_engine.pd.DataFrame(
        {
            "Open": [100 + index for index in range(25)],
            "High": [102 + index for index in range(25)],
            "Low": [99 + index for index in range(25)],
            "Close": [101 + index for index in range(25)],
            "Volume": [1000000 + index * 10000 for index in range(25)],
        },
        index=finance_engine.pd.date_range("2026-07-01", periods=25, freq="B"),
    )
    payload = finance_engine._build_finance_payload(
        "TEST",
        "stub",
        frame,
        {
            "currency": "USD",
            "pe": 24,
            "forward_pe": 20,
            "pb": 5.5,
            "market_cap": 250000000000,
            "low_52w": 80,
            "high_52w": 130,
            "fundamentals_source": "tencent_quote",
        },
    )
    payload["catalysts"] = {
        "policy": "未发现已确认的重大政策变化。",
        "earnings": "下一期财报将验证盈利兑现情况。",
        "landmark": "产品发布节奏是近期观察节点。",
        "style": "估值需结合同业与市场风格比较。",
    }
    return payload


def test_finance_analysis_is_rendered_without_changing_chart_path():
    if finance_engine.pd is None:
        return "SKIP pandas unavailable"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    finance = _finance_payload()
    original_chart_path = finance["chart_path"]
    data = [
        {
            "topic": "测试公司",
            "data": [
                {
                    "title": "测试公司发布已核验产品更新",
                    "source": "测试官方来源",
                    "date_check": "09月04日",
                    "summary": "【事件核心】\n测试公司发布产品更新。\n【深度细节/数据支撑】\n本条仅用于金融页输出验证。\n【行业深远影响】\n后续以正式披露为准。",
                    "importance": 3,
                    "chart_info": {"has_chart": False},
                }
            ],
            "report_style": "company_tracking",
            "finance": finance,
            "warnings": [],
            "extraction_stats": {},
            "focus_tags": [],
        }
    ]
    generated = Path(generate_ppt(data, [], str(OUTPUT_PPT.with_suffix("")), "stub"))
    assert generated.resolve() == OUTPUT_PPT.resolve()
    assert finance["chart_path"] == original_chart_path
    assert Path(original_chart_path).exists()

    presentation = Presentation(generated)
    finance_slide = next(
        slide
        for slide in presentation.slides
        if "量化面与事件催化" in "\n".join(
            shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
        )
    )
    text = "\n".join(
        shape.text for shape in finance_slide.shapes if getattr(shape, "has_text_frame", False)
    )
    assert "TTM市盈率: 24.00x" in text
    assert "预期市盈率: 20.00x" in text
    assert "市净率: 5.50x" in text
    assert "估值判断:" in text
    assert "股价判断:" in text
    assert "研究观点:" in text
    assert "观点依据:" in text
    assert "风险提示:" in text
    assert "行情源: stub" in text
    assert "估值源: tencent_quote" in text
    assert "不构成个性化投资建议或目标价" in text
    assert "股权风险溢价" not in text
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in finance_slide.shapes)


if __name__ == "__main__":
    test_finance_analysis_is_rendered_without_changing_chart_path()
    print(f"finance PPT output test passed: {OUTPUT_PPT}")
